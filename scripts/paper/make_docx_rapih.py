#!/usr/bin/env python3
"""Generate laporan-analisis-rapih.docx from draft-analisis-rapih.md.

Plain report style: Times New Roman, 1.5 spacing, justified body,
black headings, tables + figures in narrative order.

Usage:
  python3 scripts/paper/make_docx_rapih.py
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

REPO = Path(__file__).resolve().parents[2]
DRAFT = REPO / "latex/IEEE-Conference/draft/draft-analisis-rapih.md"
GAMBAR = REPO / "latex/IEEE-Conference/gambar"
OUT_DOCX = REPO / "latex/IEEE-Conference/draft/laporan-analisis-rapih.docx"

BLACK = RGBColor(0, 0, 0)
FONT = "Times New Roman"


def style_document(doc: Any) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    normal.font.color.rgb = BLACK
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(6)
    for i in range(1, 4):
        h = doc.styles[f"Heading {i}"]
        h.font.name = FONT
        h.font.color.rgb = BLACK
        h.font.bold = True
        h.font.size = Pt({1: 16, 2: 13, 3: 12}[i])
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.line_spacing = 1.5
        h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_body(doc: Any, text: str) -> None:
    doc.add_paragraph(text)


def add_caption(doc: Any, text: str) -> None:
    p = doc.add_paragraph("")
    run = p.add_run(text)
    run.italic = True
    run.font.name = FONT
    run.font.size = Pt(11)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_table(doc: Any, headers: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for j, h in enumerate(headers):
        c = t.cell(0, j)
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.name = FONT
        r.font.size = Pt(11)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            r = c.paragraphs[0].add_run(val)
            r.font.name = FONT
            r.font.size = Pt(10)
    doc.add_paragraph("")


def main() -> None:
    try:
        lines = DRAFT.read_text().splitlines()
    except OSError as exc:
        raise RuntimeError(f"cannot read draft: {DRAFT}: {exc}") from exc

    doc: Any = Document()
    style_document(doc)
    sec = doc.sections[0]
    sec.top_margin = Pt(72)
    sec.bottom_margin = Pt(72)
    sec.left_margin = Pt(72)
    sec.right_margin = Pt(72)

    i = 0
    pending_caption: str | None = None
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("!["):
            m = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            if m:
                caption, fname = m.group(1), m.group(2)
                pending_caption = caption
                img = GAMBAR / fname
                if img.exists():
                    doc.add_picture(str(img), width=Pt(430))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    add_body(doc, f"[Gambar hilang: {fname}]")
            else:
                add_body(doc, line)
        elif line.startswith("|"):
            tbl: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl.append(lines[i].strip())
                i += 1
            # parse: header, separator, rows
            cells = [[c.strip() for c in r.strip("|").split("|")] for r in tbl]
            if len(cells) >= 3:
                headers = cells[0]
                rows = [r for r in cells[2:] if len(r) == len(headers)]
                add_table(doc, headers, rows)
            continue
        else:
            # caption line directly after image gets italic treatment
            if pending_caption and len(line) < 300:
                add_caption(doc, line)
                pending_caption = None
            else:
                add_body(doc, line)
                pending_caption = None
        i += 1

    doc.save(str(OUT_DOCX))
    print(f"wrote {OUT_DOCX} ({OUT_DOCX.stat().st_size/1024:.0f} KB)")


if __name__ == "__main__":
    main()
